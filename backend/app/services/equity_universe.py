"""Resolve Q-FAE's stock universe against the Upstox NSE instrument master."""

from __future__ import annotations

import gzip
import hashlib
import json
import re
import threading
from dataclasses import dataclass
from datetime import UTC, datetime
from difflib import SequenceMatcher
from pathlib import Path
from typing import Any

import httpx
from openpyxl import load_workbook

UPSTOX_NSE_MASTER_URL = "https://assets.upstox.com/market-quote/instruments/exchange/NSE.json.gz"
_LEGAL_SUFFIXES = frozenset({"CO", "COMPANY", "CORP", "CORPORATION", "INDIA", "LTD", "LIMITED", "PLC"})
_NSE_EQUITY_SERIES = frozenset({"EQ", "BE", "SM", "ST", "BZ"})


@dataclass(frozen=True)
class StockListEntry:
    """A company name requested by the user-maintained stock list."""

    row_number: int
    name: str


def normalise_name(value: str) -> str:
    """Return a conservative comparison key for company and trading names."""
    words = re.findall(r"[A-Z0-9]+", value.upper())
    meaningful_words = [word for word in words if word not in _LEGAL_SUFFIXES]
    return " ".join(meaningful_words)


def load_stock_list(path: Path) -> list[StockListEntry]:
    """Read the `Name` column from the user-maintained Excel universe."""
    if not path.is_file():
        raise FileNotFoundError(f"Stock list was not found: {path}")

    workbook = load_workbook(path, read_only=True, data_only=True)
    worksheet = workbook.active
    header = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header:
        raise ValueError("Stock list has no header row")

    header_positions = {
        str(value).strip().casefold(): index
        for index, value in enumerate(header)
        if value is not None
    }
    name_index = next(
        (header_positions[key] for key in ("name", "company name", "company", "stock") if key in header_positions),
        None,
    )
    if name_index is None:
        raise ValueError("Stock list must include a Name, Company Name, Company, or Stock column")

    entries: list[StockListEntry] = []
    for row_number, row in enumerate(worksheet.iter_rows(min_row=2, values_only=True), start=2):
        raw_name = row[name_index] if name_index < len(row) else None
        if raw_name is None or not str(raw_name).strip():
            continue
        entries.append(StockListEntry(row_number=row_number, name=str(raw_name).strip()))

    if not entries:
        raise ValueError("Stock list does not contain any stock names")
    return entries


def parse_nse_equity_master(compressed_payload: bytes) -> list[dict[str, Any]]:
    """Extract NSE cash-equity series relevant to Q-FAE from the BOD master."""
    try:
        master = json.loads(gzip.decompress(compressed_payload))
    except (OSError, UnicodeDecodeError, json.JSONDecodeError) as exc:
        raise ValueError("Upstox instrument master was not a valid gzip JSON document") from exc

    if not isinstance(master, list):
        raise ValueError("Upstox instrument master has an unexpected top-level structure")

    equities = [
        instrument
        for instrument in master
        if isinstance(instrument, dict)
        and instrument.get("segment") == "NSE_EQ"
        and instrument.get("instrument_type") in _NSE_EQUITY_SERIES
        and instrument.get("instrument_key")
        and instrument.get("trading_symbol")
    ]
    if not equities:
        raise ValueError("Upstox instrument master contained no NSE equities")
    return equities


def _candidate_score(requested_name: str, instrument: dict[str, Any]) -> tuple[float, str]:
    """Find the best name-like field for a requested company name."""
    requested_key = normalise_name(requested_name)
    best_score = 0.0
    best_field = ""
    for field in ("short_name", "name", "trading_symbol"):
        value = instrument.get(field)
        if not value:
            continue
        candidate_key = normalise_name(str(value))
        if not candidate_key:
            continue
        score = SequenceMatcher(None, requested_key, candidate_key).ratio()
        if requested_key == candidate_key:
            score = 1.0
        if score > best_score:
            best_score = score
            best_field = field
    return best_score, best_field


def resolve_stock_list(
    stock_list: list[StockListEntry],
    master: list[dict[str, Any]],
    *,
    minimum_fuzzy_score: float = 0.82,
    minimum_margin: float = 0.05,
) -> list[dict[str, Any]]:
    """Resolve each requested company, retaining uncertain results for review."""
    exact_index: dict[str, list[tuple[int, str]]] = {}
    token_index: dict[str, set[int]] = {}
    for index, instrument in enumerate(master):
        for field in ("short_name", "name", "trading_symbol"):
            value = instrument.get(field)
            if not value:
                continue
            comparison_key = normalise_name(str(value))
            if not comparison_key:
                continue
            exact_index.setdefault(comparison_key, []).append((index, field))
            for token in set(comparison_key.split()):
                if len(token) >= 3:
                    token_index.setdefault(token, set()).add(index)

    resolved: list[dict[str, Any]] = []
    for entry in stock_list:
        requested_key = normalise_name(entry.name)
        exact_matches = exact_index.get(requested_key, [])
        candidate_indexes = {index for index, _ in exact_matches}
        if not candidate_indexes:
            for token in set(requested_key.split()):
                candidate_indexes.update(token_index.get(token, set()))

        if not candidate_indexes:
            resolved.append(
                {
                    "stock_list_row": entry.row_number,
                    "requested_name": entry.name,
                    "status": "unmatched",
                    "match_kind": None,
                    "match_score": 0.0,
                    "match_margin": 0.0,
                    "matched_field": None,
                    "candidate": None,
                }
            )
            continue

        ranked = sorted(
            (
                (*_candidate_score(entry.name, instrument), instrument)
                for instrument in (master[index] for index in candidate_indexes)
            ),
            key=lambda candidate: candidate[0],
            reverse=True,
        )
        best_score, matched_field, best = ranked[0]
        second_score = ranked[1][0] if len(ranked) > 1 else 0.0
        match_kind = "exact" if best_score == 1.0 else "fuzzy"
        is_safe_fuzzy_match = best_score >= minimum_fuzzy_score and best_score - second_score >= minimum_margin
        status = "resolved" if (match_kind == "exact" and len(candidate_indexes) == 1) or is_safe_fuzzy_match else "needs_review"

        row: dict[str, Any] = {
            "stock_list_row": entry.row_number,
            "requested_name": entry.name,
            "status": status,
            "match_kind": match_kind,
            "match_score": round(best_score, 4),
            "match_margin": round(best_score - second_score, 4),
            "matched_field": matched_field,
            "candidate": {
                "trading_symbol": best["trading_symbol"],
                "instrument_key": best["instrument_key"],
                "isin": best.get("isin"),
                "name": best.get("name"),
                "short_name": best.get("short_name"),
                "exchange": best.get("exchange"),
                "segment": best.get("segment"),
                "instrument_type": best.get("instrument_type"),
                "tick_size": best.get("tick_size"),
                "lot_size": best.get("lot_size"),
            },
        }
        resolved.append(row)
    return resolved


class EquityUniverseService:
    """Refresh and load the repository's Upstox-aligned equity universe."""

    def __init__(self, stock_list_path: Path, output_path: Path) -> None:
        self.stock_list_path = stock_list_path
        self.output_path = output_path
        self._refresh_lock = threading.Lock()

    def refresh(self) -> dict[str, Any]:
        """Download the current master and atomically write the resolved universe."""
        with self._refresh_lock:
            source_entries = load_stock_list(self.stock_list_path)
            compressed_master = self._download_master()
            master = parse_nse_equity_master(compressed_master)
            instruments = resolve_stock_list(source_entries, master)
            summary = {
                "requested": len(instruments),
                "resolved": sum(row["status"] == "resolved" for row in instruments),
                "needs_review": sum(row["status"] == "needs_review" for row in instruments),
                "unmatched": sum(row["status"] == "unmatched" for row in instruments),
            }
            payload = {
                "schema_version": 1,
                "generated_at": datetime.now(UTC).isoformat(),
                "source": {
                    "stock_list": self.stock_list_path.name,
                    "upstox_master_url": UPSTOX_NSE_MASTER_URL,
                    "upstox_master_sha256": hashlib.sha256(compressed_master).hexdigest(),
                },
                "summary": summary,
                "instruments": instruments,
            }
            self._write_atomically(payload)
            return payload

    def load(self) -> dict[str, Any]:
        """Load the latest generated universe without reaching Upstox."""
        if not self.output_path.is_file():
            raise FileNotFoundError("Instrument universe has not been generated yet")
        return json.loads(self.output_path.read_text(encoding="utf-8"))

    def _download_master(self) -> bytes:
        with httpx.Client(timeout=60.0, follow_redirects=True) as client:
            response = client.get(UPSTOX_NSE_MASTER_URL)
            response.raise_for_status()
            return response.content

    def _write_atomically(self, payload: dict[str, Any]) -> None:
        self.output_path.parent.mkdir(parents=True, exist_ok=True)
        temporary_path = self.output_path.with_suffix(".tmp")
        temporary_path.write_text(json.dumps(payload, indent=2, sort_keys=True) + "\n", encoding="utf-8")
        temporary_path.replace(self.output_path)
